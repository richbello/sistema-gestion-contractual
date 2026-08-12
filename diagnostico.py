#!/usr/bin/env python3
# -*- coding: utf-8 -*-
# Diagnostico: reproduce EXACTO lo que hace el navegador para el 403-2025
# Uso: python3 diagnostico.py    (desde la raiz del proyecto, con los Excel presentes)

import io, json, re, os, sys
from openpyxl import load_workbook

# localizar archivos
def buscar(*nombres):
    for n in os.listdir('.'):
        for pat in nombres:
            if pat.lower() in n.lower() and n.endswith('.xlsx'):
                return n
    return None

HIST = buscar('historico')
PLANT = buscar('formato_estadocuenta','formato')
SECOP = buscar('consolidado','secop')
print("Histórico:", HIST)
print("Plantilla:", PLANT)
print("SECOP:", SECOP)
if not (HIST and PLANT):
    print("FALTAN archivos. Necesito historico y formato_estadocuenta en esta carpeta.")
    sys.exit(1)

# Extraer el PYTHON_FILL del JS y ejecutarlo
js = open('frontend/static/js/estadocuenta-masivo.js', encoding='utf-8').read()
m = re.search(r'const PYTHON_FILL = `(.*?)`;', js, re.DOTALL)
py = m.group(1)
# ejecutar las funciones en este entorno
ns = {}
exec(py, ns)
print("\nFunciones Python del masivo cargadas OK")

# Leer histórico crudo (como hace el navegador: sin filtrar)
def leer_todo(path):
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    hdr = [str(v).strip() if v else '' for v in next(it)]
    filas = []
    for row in it:
        filas.append({hdr[i]: (row[i] if i < len(row) else None) for i in range(len(hdr))})
    wb.close()
    return filas

hist = leer_todo(HIST)
CT = '403-2025'
# EXACTO como linea 176 del JS:
pagos = [r for r in hist if CT in str(r.get('Referencia') or '')]
print(f"\nPagos crudos del {CT} (como los manda el navegador): {len(pagos)}")

if pagos:
    print("Columnas del primer pago:")
    for k in list(pagos[0].keys())[:20]:
        print(f"   '{k}'")

# Probar la deteccion tal como corre en el navegador
_detectar = ns['_detectar_cesion_masivo']
_buscar = ns['_buscar']
ces = _detectar(pagos, None)
print(f"\n_detectar_cesion_masivo devolvió: {'CESION detectada' if ces else 'None'}")
if ces:
    print(f"   cesionarios: {len(ces['cesionarios'])}")
    print(f"   valor_cesion: {ces['valor_cesion']}")
    for c in ces['cesionarios']:
        print(f"   -> {c['nombre']} ({len(c['pagos'])} pagos)")

# Probar buscar CESION en la plantilla
wb = load_workbook(PLANT)
ws = wb.active
fc = _buscar(ws, "CESION", col=2, desde=1) or _buscar(ws, "CESIÓN", col=2, desde=1)
print(f"\n_buscar CESIÓN en plantilla devolvió fila: {fc}")

# Simular el llenado completo con la funcion llenar()
import base64
pl_b64 = base64.b64encode(open(PLANT,'rb').read()).decode()
llenar = ns['llenar']
try:
    resultado_b64 = llenar(pl_b64, json.dumps(pagos, default=str), CT, '')
    out = base64.b64decode(resultado_b64)
    open('/tmp/resultado_403.xlsx','wb').write(out)
    # verificar si la cesion quedo llena
    wb2 = load_workbook('/tmp/resultado_403.xlsx')
    ws2 = wb2.active
    fc2 = _buscar(ws2, "CESION", col=2, desde=1) or _buscar(ws2, "CESIÓN", col=2, desde=1)
    fcto = _buscar(ws2, "CTO Y VIG", col=3, desde=fc2) if fc2 else None
    cesionario_llenado = ws2.cell(row=fcto+1, column=4).value if fcto else None
    print(f"\n=== RESULTADO ===")
    print(f"Bloque CESIÓN en fila: {fc2}")
    print(f"Cesionario escrito: {cesionario_llenado!r}")
    if cesionario_llenado:
        print("✅ LA CESIÓN SÍ SE LLENÓ")
    else:
        print("❌ La cesión quedó VACÍA - aquí está el problema")
    if '_DEBUG_CESION' in ns:
        print(f"DEBUG interno: {ns.get('_DEBUG_CESION')}")
except Exception as e:
    import traceback
    print(f"\n❌ ERROR en llenar(): {e}")
    traceback.print_exc()