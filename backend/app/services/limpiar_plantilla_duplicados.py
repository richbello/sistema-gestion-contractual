"""
limpiar_plantilla_duplicados.py
===============================
Post-proceso de la Plantilla de Pagos (salida del Módulo 02).

Cuando un pago se dividió en varios CRP (cruce_extraccion_crp.py), en la
plantilla aparecen bloques "duplicados": mismo contrato que el bloque anterior
pero con otro N° Interno CRP en la fila P40. En esos bloques duplicados la(s)
fila(s) P31:

  NO llevan (se dejan vacías):
    Tipo de retenciones, Indicador de retención,
    Base imponible de retención, Importe de retención
    (las retenciones se aplican una sola vez, en el primer bloque)

  SÍ llevan, copiadas del PRIMER bloque del mismo contrato:
    Código Bco, No Cuenta, Tipo Cta
    (en la Extracción la fila duplicada va sin datos bancarios y el Módulo 02
     los rellenaba con valores por defecto — p.ej. Código Bco '051' o
     No Cuenta vacío —; aquí se reemplazan por los datos reales del contratista)

Estructura del bloque (se lee como unidad, de una fila C a la siguiente):
    C   -> col 'RP Doc Presupuestal' (J) = No. contrato
    P40 -> col 'RP Doc Presupuestal' (J) = N° Interno CRP
    P31 -> datos bancarios + retención principal
    (fila de continuación opcional, sin 'Tipo Registro P', con una 2a
     retención, p.ej. 44 honorarios)

Regla para detectar el bloque duplicado:
    contrato == contrato del bloque anterior  Y  CRP (P40) != CRP del bloque anterior
    -> así NO se tocan los contratos con dos pagos sobre el mismo CRP
       (p.ej. 432-2026 y 529-2026), que son pagos completos e independientes.

Uso CLI:
    python limpiar_plantilla_duplicados.py Plantilla_Pagos.xlsx salida.xlsx

Uso en el Módulo 02 (justo después de generar la plantilla):
    from limpiar_plantilla_duplicados import limpiar_bloques_duplicados
    resumen = limpiar_bloques_duplicados(ruta_plantilla, ruta_plantilla)
"""

from __future__ import annotations

import sys
import unicodedata

import openpyxl

# Columnas que se dejan VACÍAS en las filas P31 de los bloques duplicados
COLS_VACIAR_EN_DUPLICADO = [
    "Tipo de retenciones",
    "Indicador de retención",
    "Base imponible de retención",
    "Importe de retención",
]

# Columnas que en los bloques duplicados se copian del primer bloque del contrato
COLS_COPIAR_DEL_PRIMERO = [
    "Código Bco",
    "No Cuenta",
    "Tipo Cta",
]

H_TIPO_REG = "Tipo Registro P"
H_CLAVE = "Clave Contab."
H_RP = "RP Doc Presupuestal"


def _norm(s) -> str:
    if s is None:
        return ""
    s = unicodedata.normalize("NFKD", str(s))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return " ".join(s.lower().split())


def _col(headers, nombre):
    """Índice 1-based de la ÚLTIMA columna con ese encabezado (las de banco y
    retención están al final; 'Tipo Ident' se repite, pero no se usa aquí)."""
    objetivo = _norm(nombre)
    idx = [i + 1 for i, h in enumerate(headers) if _norm(h) == objetivo]
    if not idx:
        raise ValueError(f"No se encontró la columna '{nombre}' en la plantilla")
    return idx[-1]


def _clave(v):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


def _leer_bloques(ws, c_tipo, c_clave, c_rp):
    """Lista de bloques: dict(fila_c, filas=[...], contrato, crp, filas_p31)."""
    bloques = []
    actual = None
    for r in range(2, ws.max_row + 1):
        tipo = str(ws.cell(r, c_tipo).value or "").strip().upper()
        if tipo == "C":
            actual = {"fila_c": r, "filas": [r],
                      "contrato": str(ws.cell(r, c_rp).value or "").strip(),
                      "crp": None, "filas_p31": [], "filas_cont": []}
            bloques.append(actual)
            continue
        if actual is None:
            continue
        actual["filas"].append(r)
        clave = _clave(ws.cell(r, c_clave).value)
        if tipo == "P" and clave == 40 and actual["crp"] is None:
            actual["crp"] = str(ws.cell(r, c_rp).value or "").strip()
        elif tipo == "P" and clave == 31:
            actual["filas_p31"].append(r)
        elif tipo == "" and any(ws.cell(r, c).value not in (None, "") for c in range(1, ws.max_column + 1)):
            # fila de continuación (2a retención) del P31 anterior
            actual["filas_cont"].append(r)
    return bloques


def limpiar_bloques_duplicados(ruta_plantilla: str, ruta_salida: str) -> dict:
    wb = openpyxl.load_workbook(ruta_plantilla)
    ws = wb.worksheets[0]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]

    c_tipo = _col(headers, H_TIPO_REG)
    c_clave = _col(headers, H_CLAVE)
    c_rp = headers.index(next(h for h in headers if _norm(h) == _norm(H_RP))) + 1
    cols_vaciar = [_col(headers, n) for n in COLS_VACIAR_EN_DUPLICADO]
    cols_copiar = [_col(headers, n) for n in COLS_COPIAR_DEL_PRIMERO]

    bloques = _leer_bloques(ws, c_tipo, c_clave, c_rp)

    duplicados = []
    filas_a_borrar = []
    primero = bloques[0] if bloques else None   # primer bloque del contrato actual
    for prev, b in zip(bloques, bloques[1:]):
        if b["contrato"] and b["contrato"] == prev["contrato"] and b["crp"] != prev["crp"]:
            # Se compara contra el bloque anterior; en una división en 3 CRP
            # el 3er bloque también cumple (mismo contrato, otro CRP).
            banco = {}
            if primero["filas_p31"]:
                r0 = primero["filas_p31"][0]
                banco = {c: ws.cell(r0, c).value for c in cols_copiar}
            for r in b["filas_p31"]:
                for c in cols_vaciar:
                    ws.cell(r, c).value = None
                for c, v in banco.items():
                    ws.cell(r, c).value = v
            # Una fila de continuación solo trae retención: en el duplicado
            # quedaría vacía, así que se elimina para no dejar filas en blanco.
            filas_a_borrar.extend(b["filas_cont"])
            duplicados.append(f"{b['contrato']} (CRP {b['crp']}, fila {b['fila_c']})")
        else:
            primero = b

    for r in sorted(filas_a_borrar, reverse=True):
        ws.delete_rows(r)

    wb.save(ruta_salida)
    return {"bloques": len(bloques), "bloques_duplicados_limpiados": len(duplicados),
            "detalle": duplicados}


if __name__ == "__main__":
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    res = limpiar_bloques_duplicados(sys.argv[1], sys.argv[2])
    print(f"bloques: {res['bloques']}")
    print(f"bloques duplicados limpiados: {res['bloques_duplicados_limpiados']}")
    for d in res["detalle"]:
        print("  -", d)
