"""
cruce_extraccion_crp.py
=======================
Cruce Extracción de pagos  ↔  Reporte CRP (BogData)

Para cada fila de la Extracción:
  1. Toma el Contrato (col A) y lo busca en el Reporte CRP por 'No. Compromiso'.
  2. Toma solo los CRP con 'Com.Sin.Aut.Giro' > 0 (saldo disponible para girar).
  3. Los ordena por 'Fecha Registro' (más antiguo primero) y consume su saldo
     para cubrir el 'Valor Bruto' de la fila.
  4. Si un CRP no alcanza, la fila se divide: una fila por cada CRP usado,
     cada una con su N° Interno CRP y la parte del valor que cubre.
     La 1a fila conserva descuentos y cuenta bancaria; las duplicadas van
     sin BASE RETEICA, Valor Reteica, TOTAL DESCUENTOS, HONOR_*, Código Bco,
     Tipo Cta ni No Cuenta.
  5. El saldo consumido se descuenta, de modo que si el mismo contrato aparece
     dos veces en la Extracción (dos pagos), el segundo pago usa el saldo que
     quedó después del primero.

Salida:
  - Hoja 1 con la misma estructura A:T; la columna 'CRP' pasa a llamarse
    'N° Interno CRP' y lleva el CRP asignado. Montos como número.
  - Hoja 'Cruce_CRP' con el detalle del cruce (auditoría).

Las columnas se ubican por NOMBRE de encabezado, no por posición, porque el
layout del Reporte CRP cambia entre reportes.

Uso CLI:
    python cruce_extraccion_crp.py Extraccion.xlsx Reporte_CRP.xlsx salida.xlsx

Uso como servicio (Flask/FastAPI):
    from cruce_extraccion_crp import cruzar_extraccion_con_crp
    resumen = cruzar_extraccion_con_crp(ruta_extr, ruta_crp, ruta_salida)
"""

from __future__ import annotations

import sys
import unicodedata
from copy import copy
from dataclasses import dataclass, field
from datetime import datetime, date
from decimal import Decimal, ROUND_HALF_UP

import openpyxl
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# CONFIGURACIÓN
# ---------------------------------------------------------------------------

# Encabezados en la Extracción
EXT_CONTRATO = "Contrato"
EXT_VALOR_BRUTO = "Valor Bruto"

# Columnas que SOLO van en la primera fila cuando un pago se divide en varios
# CRP. En las filas duplicadas (2a, 3a...) quedan VACÍAS: los descuentos y la
# cuenta bancaria se aplican una sola vez, en la primera fila.
EXT_COLS_SOLO_PRIMERA = [
    "BASE RETEICA",
    "Valor Reteica",
    "TOTAL DESCUENTOS",
    "HONOR_BASE",
    "HONOR_VALOR",
    "Código Bco",
    "Tipo Cta",
    "No Cuenta",
]

# Neto a Pagar se recalcula por fila:
#   1a fila  = Neto original - (Valor Bruto original - Valor Bruto de la fila)
#              (es decir, su parte del bruto menos TODOS los descuentos)
#   duplicadas = su parte del Valor Bruto (no llevan descuentos)
# Así la suma de las filas sigue dando exactamente el Neto original.
EXT_NETO = "Neto a Pagar"

# El N° Interno CRP asignado se escribe EN la columna 'CRP' de la Extracción,
# que pasa a llamarse 'N° Interno CRP' (misma posición, A:T no cambia).
# Si la Extracción no trae columna 'CRP', se agrega al final.
EXT_CRP = "CRP"
EXT_COL_NUEVA = "N° Interno CRP"

# Columnas que en la salida quedan como NÚMERO (Módulo 01 las deja como texto).
# Código Bco, Tipo Cta, No Cuenta y N° Interno CRP siguen como texto (ceros a
# la izquierda).
EXT_COLS_NUMERICAS = [
    "NIT_CC", "Valor Bruto", "BASE RETEICA", "Valor Reteica", "TOTAL DESCUENTOS",
    "Neto a Pagar", "PERIODO", "PAGO No.", "HONOR_BASE", "HONOR_VALOR", "SALDO CRP",
]

# Encabezados en el Reporte CRP
CRP_COMPROMISO = "No. Compromiso"
CRP_INTERNO = "N° Interno CRP"
CRP_SALDO = "Com.Sin.Aut.Giro"
CRP_FECHA = "Fecha Registro"

# Marcas cuando no hay CRP
MARCA_SIN_CRUCE = "SIN CRUCE"
MARCA_SIN_SALDO = "SALDO INSUFICIENTE"

FILL_DIVIDIDA = PatternFill("solid", fgColor="FFF2CC")  # amarillo: fila dividida
FILL_ERROR = PatternFill("solid", fgColor="F8CBAD")     # rojo: sin cruce / sin saldo
FILL_HEADER_NUEVO = PatternFill("solid", fgColor="C6E0B4")


# ---------------------------------------------------------------------------
# NORMALIZACIÓN
# ---------------------------------------------------------------------------

def _norm_header(s) -> str:
    """Encabezado sin tildes, sin espacios dobles, en minúscula."""
    if s is None:
        return ""
    s = unicodedata.normalize("NFKD", str(s))
    s = "".join(c for c in s if not unicodedata.combining(c))
    return " ".join(s.lower().split())


def norm_compromiso(v) -> str:
    """'251 2026' / '251-2026' / ' 251-2026 ' -> '251-2026'."""
    if v is None:
        return ""
    s = str(v).strip().upper()
    s = "-".join(s.split())          # espacios -> guion
    while "--" in s:
        s = s.replace("--", "-")
    return s


def to_decimal(v) -> Decimal:
    """Convierte '7839000', 7839000, '7.839.000', '7,839,000.00' a Decimal."""
    if v is None or v == "":
        return Decimal(0)
    if isinstance(v, (int, float, Decimal)):
        return Decimal(str(v))
    s = str(v).strip().replace("$", "").replace(" ", "")
    if not s:
        return Decimal(0)
    # Formato colombiano 7.839.000,50
    if "," in s and "." in s:
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif s.count(".") > 1:
        s = s.replace(".", "")
    elif "," in s:
        s = s.replace(",", ".") if s.count(",") == 1 and len(s.split(",")[1]) <= 2 else s.replace(",", "")
    try:
        return Decimal(s)
    except Exception:
        return Decimal(0)


def to_date(v):
    if isinstance(v, datetime):
        return v
    if isinstance(v, date):
        return datetime(v.year, v.month, v.day)
    if v is None:
        return datetime.max
    s = str(v).strip()[:10]
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            pass
    return datetime.max


def _escribir_como_original(original, nuevo: Decimal):
    """Devuelve el valor nuevo con el mismo TIPO que tenía la celda original
    (la Extracción guarda los montos como texto; se respeta para no romper
    el cargue posterior en Plantilla de pagos)."""
    entero = nuevo.quantize(Decimal(1), rounding=ROUND_HALF_UP)
    if isinstance(original, str):
        return str(entero) if nuevo == entero else str(nuevo)
    if isinstance(original, float):
        return float(nuevo)
    return int(entero) if nuevo == entero else float(nuevo)


def _a_numero(v):
    """'4703400' -> 4703400, '0.766' -> 0.766; lo que no sea número queda igual."""
    if v is None or isinstance(v, (int, float)):
        return v
    s = str(v).strip()
    if not s:
        return None
    try:
        d = Decimal(s)
    except Exception:
        return v
    return int(d) if d == d.to_integral_value() else float(d)


def _buscar_columnas(headers, requeridas, origen):
    idx = {_norm_header(h): i for i, h in enumerate(headers)}
    res = {}
    faltan = []
    for nombre in requeridas:
        i = idx.get(_norm_header(nombre))
        if i is None:
            faltan.append(nombre)
        else:
            res[nombre] = i
    if faltan:
        raise ValueError(f"[{origen}] No se encontraron las columnas: {faltan}")
    return res


# ---------------------------------------------------------------------------
# CARGA DEL REPORTE CRP
# ---------------------------------------------------------------------------

@dataclass
class CRPDisponible:
    interno: str
    fecha: datetime
    saldo_inicial: Decimal
    saldo: Decimal = field(default=Decimal(0))

    def __post_init__(self):
        self.saldo = self.saldo_inicial


def cargar_crp(ruta_crp: str, hoja: str | None = None) -> dict[str, list[CRPDisponible]]:
    """Devuelve {compromiso_normalizado: [CRPDisponible ordenados por fecha]}.
    Solo incluye CRP con Com.Sin.Aut.Giro > 0."""
    wb = openpyxl.load_workbook(ruta_crp, read_only=True, data_only=True)
    ws = wb[hoja] if hoja else wb.worksheets[0]
    filas = ws.iter_rows(values_only=True)
    headers = next(filas)
    col = _buscar_columnas(headers, [CRP_COMPROMISO, CRP_INTERNO, CRP_SALDO, CRP_FECHA], "Reporte CRP")

    base: dict[str, list[CRPDisponible]] = {}
    vistos = set()
    for r in filas:
        if r is None:
            continue
        comp = norm_compromiso(r[col[CRP_COMPROMISO]])
        if not comp:
            continue
        saldo = to_decimal(r[col[CRP_SALDO]])
        if saldo <= 0:
            continue
        interno = str(r[col[CRP_INTERNO]] or "").strip()
        if interno.endswith(".0"):
            interno = interno[:-2]
        # Un CRP puede tener varias posiciones; si el mismo N° Interno aparece
        # varias veces para el mismo compromiso se suman sus saldos.
        clave = (comp, interno)
        if clave in vistos:
            for c in base[comp]:
                if c.interno == interno:
                    c.saldo_inicial += saldo
                    c.saldo += saldo
            continue
        vistos.add(clave)
        base.setdefault(comp, []).append(
            CRPDisponible(interno=interno, fecha=to_date(r[col[CRP_FECHA]]), saldo_inicial=saldo)
        )
    wb.close()

    for lst in base.values():
        lst.sort(key=lambda c: (c.fecha, c.interno))
    return base


# ---------------------------------------------------------------------------
# REPARTO DE UNA FILA ENTRE VARIOS CRP
# ---------------------------------------------------------------------------

def asignar_crp(valor: Decimal, disponibles: list[CRPDisponible]):
    """Consume saldo de los CRP (más antiguo primero) hasta cubrir 'valor'.
    Devuelve lista de (interno, monto_asignado, saldo_antes) y el faltante."""
    asignaciones = []
    pendiente = valor
    for c in disponibles:
        if pendiente <= 0:
            break
        if c.saldo <= 0:
            continue
        toma = min(c.saldo, pendiente)
        asignaciones.append((c.interno, toma, c.saldo))
        c.saldo -= toma
        pendiente -= toma
    return asignaciones, pendiente


# ---------------------------------------------------------------------------
# PROCESO PRINCIPAL
# ---------------------------------------------------------------------------

def cruzar_extraccion_con_crp(ruta_extraccion: str, ruta_crp: str, ruta_salida: str,
                              conservar_sin_cruce: bool = False,
                              actualizar_saldo: bool = False) -> dict:
    """conservar_sin_cruce: si el contrato no aparece en el Reporte CRP, la fila se deja
    con el CRP que ya traía (SIPAGO lo trae por cédula) en vez de marcar SIN CRUCE.
    actualizar_saldo: la columna SALDO CRP pasa a ser el saldo del CRP asignado."""
    base = cargar_crp(ruta_crp)

    wb_in = openpyxl.load_workbook(ruta_extraccion)
    ws_in = wb_in.worksheets[0]
    max_col = ws_in.max_column  # A:T = 20
    headers = [ws_in.cell(1, c).value for c in range(1, max_col + 1)]
    col = _buscar_columnas(headers, [EXT_CONTRATO, EXT_VALOR_BRUTO], "Extracción")
    cols_solo_primera = []
    for n in EXT_COLS_SOLO_PRIMERA:
        try:
            cols_solo_primera.append(_buscar_columnas(headers, [n], "Extracción")[n])
        except ValueError:
            pass  # columna opcional
    try:
        col_neto = _buscar_columnas(headers, [EXT_NETO], "Extracción")[EXT_NETO]
    except ValueError:
        col_neto = None
    cols_numericas = set()
    for n in EXT_COLS_NUMERICAS:
        try:
            cols_numericas.add(_buscar_columnas(headers, [n], "Extracción")[n])
        except ValueError:
            pass
    # Columna donde va el N° Interno CRP: la propia 'CRP' (o la ya renombrada)
    col_saldo = next((i for i, hh in enumerate(headers)
                      if _norm_header(hh) == _norm_header("SALDO CRP")), None)
    col_crp = next((i for i, hh in enumerate(headers)
                    if _norm_header(hh) in (_norm_header(EXT_CRP), _norm_header(EXT_COL_NUEVA))), None)

    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = ws_in.title

    for c in range(1, max_col + 1):
        src = ws_in.cell(1, c)
        dst = ws_out.cell(1, c, src.value)
        _copiar_estilo(src, dst)
    if col_crp is not None:
        col_nueva = col_crp + 1
        ws_out.cell(1, col_nueva, EXT_COL_NUEVA)
        ancho_total = max_col
    else:
        col_nueva = max_col + 1
        h = ws_out.cell(1, col_nueva, EXT_COL_NUEVA)
        _copiar_estilo(ws_in.cell(1, max_col), h)
        h.font = Font(bold=True)
        ancho_total = col_nueva

    for letra, dim in ws_in.column_dimensions.items():
        ws_out.column_dimensions[letra].width = dim.width
    ws_out.column_dimensions[get_column_letter(col_nueva)].width = 16
    ws_out.freeze_panes = ws_in.freeze_panes

    log = []
    resumen = {"filas_entrada": 0, "filas_salida": 0, "filas_divididas": 0,
               "sin_cruce": 0, "saldo_insuficiente": 0}
    fila_out = 2

    for r in range(2, ws_in.max_row + 1):
        celdas = [ws_in.cell(r, c) for c in range(1, max_col + 1)]
        valores = [c.value for c in celdas]
        if all(v in (None, "") for v in valores):
            continue
        resumen["filas_entrada"] += 1

        contrato = norm_compromiso(valores[col[EXT_CONTRATO]])
        bruto = to_decimal(valores[col[EXT_VALOR_BRUTO]])
        disponibles = base.get(contrato, [])

        if not disponibles:
            asign, faltante = [], bruto
        else:
            asign, faltante = asignar_crp(bruto, disponibles)

        # Construir la lista de (interno, monto) para las filas de salida
        partes = [(i, m) for i, m, _ in asign]
        crp_previo = valores[col_crp] if col_crp is not None else None
        if not disponibles and conservar_sin_cruce and crp_previo not in (None, ""):
            partes = [(str(crp_previo), bruto)]
            faltante = Decimal(0)
            resumen["sin_cruce"] += 1
            resumen.setdefault("conservados", []).append(str(valores[col[EXT_CONTRATO]]))
        if faltante > 0:
            partes.append((MARCA_SIN_CRUCE if not disponibles else MARCA_SIN_SALDO, faltante))
            resumen["sin_cruce" if not disponibles else "saldo_insuficiente"] += 1

        neto_orig = to_decimal(valores[col_neto]) if col_neto is not None else None

        if len(partes) > 1:
            resumen["filas_divididas"] += 1

        for k, (interno, monto) in enumerate(partes):
            nuevos = list(valores)
            if len(partes) > 1:
                nuevos[col[EXT_VALOR_BRUTO]] = _escribir_como_original(valores[col[EXT_VALOR_BRUTO]], monto)
                if k == 0:
                    neto = neto_orig - (bruto - monto) if neto_orig is not None else None
                else:
                    neto = monto
                    for i in cols_solo_primera:
                        nuevos[i] = None
                if col_neto is not None:
                    if neto < 0:
                        resumen.setdefault("alertas", []).append(
                            f"Fila {r} ({valores[col[EXT_CONTRATO]]}): la 1a parte no cubre los descuentos")
                    nuevos[col_neto] = _escribir_como_original(valores[col_neto], neto)
            for c in range(max_col):
                v = _a_numero(nuevos[c]) if c in cols_numericas else nuevos[c]
                dst = ws_out.cell(fila_out, c + 1, v)
                _copiar_estilo(celdas[c], dst)
                if c in cols_numericas:
                    dst.number_format = "General"
            ws_out.cell(fila_out, col_nueva, interno).number_format = "@"
            if actualizar_saldo and col_saldo is not None and disponibles:
                s_ant = next((s for i, _, s in asign if i == interno), None)
                ws_out.cell(fila_out, col_saldo + 1, float(s_ant) if s_ant is not None else None)
            if interno in (MARCA_SIN_CRUCE, MARCA_SIN_SALDO):
                for c in range(1, ancho_total + 1):
                    ws_out.cell(fila_out, c).fill = FILL_ERROR
            elif len(partes) > 1:
                for c in range(1, ancho_total + 1):
                    ws_out.cell(fila_out, c).fill = FILL_DIVIDIDA

            saldo_antes = next((s for i, _, s in asign if i == interno), None)
            log.append({
                "Fila origen": r,
                "Contrato": valores[col[EXT_CONTRATO]],
                "Valor Bruto original": float(bruto),
                "Parte": f"{k + 1} de {len(partes)}",
                "N° Interno CRP": interno,
                "Saldo CRP antes": float(saldo_antes) if saldo_antes is not None else None,
                "Valor asignado": float(monto),
                "Saldo CRP después": float(saldo_antes - monto) if saldo_antes is not None else None,
                "CRP disponibles": ", ".join(f"{c.interno} ({c.fecha:%d/%m/%Y}: {c.saldo_inicial:,.0f})"
                                             for c in disponibles),
            })
            fila_out += 1

    resumen["filas_salida"] = fila_out - 2
    _hoja_log(wb_out, log)
    wb_out.save(ruta_salida)
    return resumen


def _copiar_estilo(src, dst):
    if src.has_style:
        dst.font = copy(src.font)
        dst.border = copy(src.border)
        dst.fill = copy(src.fill)
        dst.number_format = src.number_format
        dst.alignment = copy(src.alignment)


def _hoja_log(wb, log):
    ws = wb.create_sheet("Cruce_CRP")
    if not log:
        return
    cab = list(log[0].keys())
    ws.append(cab)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = FILL_HEADER_NUEVO
    for fila in log:
        ws.append([fila[k] for k in cab])
        if fila["N° Interno CRP"] in (MARCA_SIN_CRUCE, MARCA_SIN_SALDO):
            for c in ws[ws.max_row]:
                c.fill = FILL_ERROR
    for i, w in enumerate([10, 12, 18, 8, 16, 16, 16, 16, 70], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for col_letra in "CFGH":
        for c in ws[col_letra][1:]:
            c.number_format = "#,##0"
    ws.freeze_panes = "A2"


# ---------------------------------------------------------------------------
if __name__ == "__main__":
    if len(sys.argv) != 4:
        print(__doc__)
        sys.exit(1)
    res = cruzar_extraccion_con_crp(sys.argv[1], sys.argv[2], sys.argv[3])
    for k, v in res.items():
        print(f"{k}: {v}")
