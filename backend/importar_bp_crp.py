# -*- coding: utf-8 -*-
"""Importa Validacion_BP_saldoCRP.xlsx -> app/data/bp_crp.json"""
import sys, os, re, json
from datetime import datetime
import openpyxl


def digitos(v):
    return re.sub(r"\D", "", str(v or ""))


def norm(v):
    return re.sub(r"\s+", " ", str(v or "").strip().upper())


def main():
    ruta = sys.argv[1] if len(sys.argv) > 1 else "Validacion_BP_saldoCRP.xlsx"
    wb = openpyxl.load_workbook(ruta, data_only=True)
    ws = wb["Hoja1"] if "Hoja1" in wb.sheetnames else wb.active

    documentos_bp = {}
    lineas = []

    for r in range(2, ws.max_row + 1):
        contrato = norm(ws.cell(r, 1).value)
        bp = digitos(ws.cell(r, 2).value)
        doc = digitos(ws.cell(r, 4).value)
        nombre = norm(ws.cell(r, 5).value)
        crp = digitos(ws.cell(r, 12).value)
        saldo = ws.cell(r, 11).value or 0
        try:
            saldo = float(saldo)
        except (ValueError, TypeError):
            saldo = 0.0
        if not doc:
            continue
        if doc not in documentos_bp:
            documentos_bp[doc] = {"nombre": nombre, "bp": bp}
        lineas.append({"contrato": contrato, "doc": doc, "crp": crp,
                       "saldo": saldo, "nombre": nombre})

    salida = {
        "meta": {
            "generado": datetime.now().isoformat(timespec="seconds"),
            "fuente": os.path.basename(ruta),
            "documentos_con_bp": len(documentos_bp),
            "lineas": len(lineas),
        },
        "documentos_bp": documentos_bp,
        "lineas": lineas,
    }

    destino = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                           "app", "data", "bp_crp.json")
    os.makedirs(os.path.dirname(destino), exist_ok=True)
    with open(destino, "w", encoding="utf-8") as f:
        json.dump(salida, f, ensure_ascii=False, indent=2)

    print("Documentos con BP:", len(documentos_bp))
    print("Lineas CRP:", len(lineas))
    print("Archivo generado:", destino)


if __name__ == "__main__":
    main()