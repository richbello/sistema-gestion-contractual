# -*- coding: utf-8 -*-
"""
Genera un Excel con los contratistas de la extraccion que NO estan en el
registro depurado. Marca los casos donde el NOMBRE ya existe en el registro
pero con otra cedula (posible cedula mal leida en el PDF).
"""
import json, re
import openpyxl
from openpyxl.styles import Font, Alignment

REG = "app/data/cuentas_bancarias.json"
PLANTILLA = "plantilla_prueba.xlsx"
SALIDA = "contratistas_faltantes.xlsx"


def norm(s):
    return re.sub(r"\s+", " ", str(s or "").strip().upper())


reg = json.load(open(REG, encoding="utf-8"))["cuentas"]
por_nombre = {}
for ced, r in reg.items():
    por_nombre.setdefault(norm(r["nombre"]), []).append(ced)

wb = openpyxl.load_workbook(PLANTILLA)
ws = wb.active
filas = list(ws.iter_rows(min_row=2))

faltantes = {}
i = 0
while i < len(filas):
    r = filas[i]
    if r[0].value == "C":
        nombre = norm(r[10].value or "")
        ced = banco = cuenta = tipo = ""
        if i + 2 < len(filas) and filas[i + 2][1].value == 31:
            p31 = filas[i + 2]
            ced = str(p31[4].value or "")
            banco = str(p31[36].value or "")
            cuenta = str(p31[37].value or "")
            tipo = str(p31[38].value or "")
        if ced and ced not in reg and ced not in faltantes:
            obs = ""
            if nombre in por_nombre:
                obs = "VERIFICAR CEDULA: nombre existe en registro con cedula " \
                      + "/".join(por_nombre[nombre])
            faltantes[ced] = {"nombre": nombre, "tipo": tipo,
                              "banco": banco, "cuenta": cuenta, "obs": obs}
        i += 3
    else:
        i += 1

out = openpyxl.Workbook()
o = out.active
o.title = "Faltantes"
headers = ["Nº identificación", "Nombre", "Tipo de Cuenta",
           "Codigo de Banco", "Numero Cuenta Bancaria", "Observación"]
for c, h in enumerate(headers, 1):
    o.cell(1, c, h).font = Font(bold=True)

fila = 2
for ced, d in sorted(faltantes.items(), key=lambda x: x[1]["nombre"]):
    o.cell(fila, 1, ced).number_format = "@"
    o.cell(fila, 2, d["nombre"])
    o.cell(fila, 3, d["tipo"]).number_format = "@"
    o.cell(fila, 4, d["banco"]).number_format = "@"
    o.cell(fila, 5, d["cuenta"]).number_format = "@"
    o.cell(fila, 6, d["obs"])
    fila += 1

anchos = [18, 34, 14, 16, 24, 55]
for idx, w in enumerate(anchos, 1):
    o.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = w
for row in o.iter_rows(min_row=1):
    for cell in row:
        cell.alignment = Alignment(horizontal="left")

out.save(SALIDA)

marcados = sum(1 for d in faltantes.values() if d["obs"])
print("=" * 60)
print(f"  Contratistas faltantes (unicos) : {len(faltantes)}")
print(f"  Marcados 'VERIFICAR CEDULA'      : {marcados}")
print(f"  Archivo generado                 : backend/{SALIDA}")
print("=" * 60)
for ced, d in list(sorted(faltantes.items(), key=lambda x: x[1]['nombre']))[:8]:
    print(f"  {ced:<12} | {d['nombre'][:30]:<30} | {d['banco']} | {d['cuenta']}")
