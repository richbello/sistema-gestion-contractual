#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generador de informes PDF y Excel para Vigilancia Fiscal IA"""
import io
from datetime import datetime
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                TableStyle, PageBreak)
from reportlab.lib.enums import TA_CENTER, TA_LEFT
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# Paleta institucional (azul-indigo, igual que el sidebar)
AZUL_OSCURO = colors.HexColor("#0F172A")
AZUL = colors.HexColor("#1E3A5F")
AZUL_OCEANO = colors.HexColor("#0C4A6E")
GRIS = colors.HexColor("#64748B")
GRIS_CLARO = colors.HexColor("#F1F5F9")
NIVEL_COLOR = {
    "CRITICA": colors.HexColor("#DC2626"),
    "ALTA": colors.HexColor("#EA580C"),
    "MEDIA": colors.HexColor("#CA8A04"),
    "INFO": colors.HexColor("#64748B"),
}

def _fmt_cop(v):
    try:
        n = int(float(str(v).replace(",", "").split(".")[0] or 0))
        return "$" + f"{n:,}".replace(",", ".")
    except Exception:
        return "$0"

def generar_pdf(resultado):
    """Genera un PDF en memoria a partir del resultado del analisis."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=2*cm, bottomMargin=2*cm,
                            leftMargin=2*cm, rightMargin=2*cm,
                            title=f"Vigilancia Fiscal - {resultado.get('entidad','')}")
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle("Titulo", parent=styles["Title"], textColor=AZUL_OSCURO, fontSize=22, spaceAfter=6))
    styles.add(ParagraphStyle("Sub", parent=styles["Normal"], textColor=GRIS, fontSize=11, alignment=TA_CENTER, spaceAfter=2))
    styles.add(ParagraphStyle("H2", parent=styles["Heading2"], textColor=AZUL_OCEANO, fontSize=14, spaceBefore=14, spaceAfter=8))
    styles.add(ParagraphStyle("Cuerpo", parent=styles["Normal"], fontSize=10, leading=15, spaceAfter=6))
    styles.add(ParagraphStyle("Pie", parent=styles["Normal"], fontSize=8, textColor=GRIS))

    elems = []

    # ---- PORTADA ----
    elems.append(Spacer(1, 1.5*cm))
    elems.append(Paragraph("VIGILANCIA FISCAL IA", styles["Titulo"]))
    elems.append(Paragraph("Cruce SECOP II &times; SIRI (Procuradur&iacute;a) &times; Bolet&iacute;n (Contralor&iacute;a)", styles["Sub"]))
    elems.append(Spacer(1, 0.8*cm))

    meta = [
        ["Entidad auditada:", resultado.get("entidad", "")],
        ["Fecha de an\u00e1lisis:", resultado.get("fecha_analisis", "")],
        ["Contratos analizados:", str(resultado.get("total_contratos", 0))],
        ["Registros SIRI:", f"{resultado.get('registros_siri', 0):,}".replace(",", ".")],
        ["Registros Bolet\u00edn:", f"{resultado.get('registros_boletin', 0):,}".replace(",", ".")],
    ]
    t = Table(meta, colWidths=[5*cm, 10*cm])
    t.setStyle(TableStyle([
        ("FONTNAME", (0,0), (0,-1), "Helvetica-Bold"),
        ("TEXTCOLOR", (0,0), (0,-1), AZUL_OSCURO),
        ("FONTSIZE", (0,0), (-1,-1), 10),
        ("BOTTOMPADDING", (0,0), (-1,-1), 6),
        ("LINEBELOW", (0,0), (-1,-1), 0.5, GRIS_CLARO),
    ]))
    elems.append(t)
    elems.append(Spacer(1, 0.8*cm))

    # ---- RESUMEN EJECUTIVO ----
    elems.append(Paragraph("Resumen ejecutivo", styles["H2"]))
    resumen = [
        ["Nivel", "Cantidad", "Significado"],
        ["CR\u00cdTICA", str(resultado.get("criticas", 0)), "Contrato firmado DESPU\u00c9S de la inhabilidad"],
        ["ALTA", str(resultado.get("altas", 0)), "Inhabilidad vigente (vigilar renovaci\u00f3n)"],
        ["MEDIA", str(resultado.get("medias", 0)), "Antecedente con inhabilidad vencida"],
        ["INFO", str(resultado.get("info", 0)), "Antecedente sin inhabilidad vigente"],
        ["TOTAL", str(resultado.get("total_alertas", 0)), ""],
    ]
    tr = Table(resumen, colWidths=[3*cm, 2.5*cm, 9.5*cm])
    estilo_r = [
        ("BACKGROUND", (0,0), (-1,0), AZUL_OSCURO),
        ("TEXTCOLOR", (0,0), (-1,0), colors.white),
        ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
        ("FONTNAME", (0,-1), (-1,-1), "Helvetica-Bold"),
        ("FONTSIZE", (0,0), (-1,-1), 9),
        ("GRID", (0,0), (-1,-1), 0.5, GRIS_CLARO),
        ("ROWBACKGROUNDS", (0,1), (-1,-2), [colors.white, GRIS_CLARO]),
        ("VALIGN", (0,0), (-1,-1), "MIDDLE"),
        ("TOPPADDING", (0,0), (-1,-1), 5),
        ("BOTTOMPADDING", (0,0), (-1,-1), 5),
    ]
    for i, nivel in enumerate(["CRITICA","ALTA","MEDIA","INFO"], start=1):
        estilo_r.append(("TEXTCOLOR", (0,i), (0,i), NIVEL_COLOR[nivel]))
        estilo_r.append(("FONTNAME", (0,i), (0,i), "Helvetica-Bold"))
    tr.setStyle(TableStyle(estilo_r))
    elems.append(tr)

    # ---- METODOLOGIA ----
    elems.append(Paragraph("Metodolog\u00eda y Regla de Oro", styles["H2"]))
    elems.append(Paragraph(
        "Este informe cruza los contratos publicados en el SECOP II de la entidad auditada "
        "contra dos fuentes oficiales: el <b>SIRI</b> (Sistema de Informaci&oacute;n de Registros y Sanciones "
        "de la Procuradur&iacute;a General de la Naci&oacute;n) y el <b>Bolet&iacute;n de Responsables Fiscales</b> "
        "de la Contralor&iacute;a General de la Rep&uacute;blica. El cruce se realiza por n&uacute;mero de "
        "documento del proveedor adjudicado.", styles["Cuerpo"]))
    elems.append(Paragraph(
        "<b>Regla de Oro (criterio temporal):</b> una alerta se clasifica como <b>CR&Iacute;TICA</b> "
        "&uacute;nicamente cuando el contrato fue firmado en una fecha <b>posterior</b> a la fecha en que "
        "la persona qued&oacute; inhabilitada y dicha inhabilidad segu&iacute;a vigente. Si la inhabilidad "
        "est&aacute; vigente pero el contrato es anterior, se marca <b>ALTA</b> (vigilar renovaciones). "
        "Los antecedentes sin inhabilidad vigente se clasifican como MEDIA o INFO.", styles["Cuerpo"]))
    elems.append(Paragraph(
        "Cada coincidencia debe ser verificada por el analista antes de cualquier actuaci&oacute;n, "
        "descartando homonimias. El dato se presenta como insumo de control; la interpretaci&oacute;n "
        "jur&iacute;dica corresponde al ente competente.", styles["Cuerpo"]))

    elems.append(PageBreak())

    # ---- TABLA DE ALERTAS ----
    elems.append(Paragraph(f"Detalle de alertas ({resultado.get('total_alertas',0)})", styles["H2"]))
    alertas = resultado.get("alertas", [])
    if not alertas:
        elems.append(Paragraph("No se encontraron coincidencias para esta entidad.", styles["Cuerpo"]))
    else:
        encab = ["Nivel", "Proveedor / Doc.", "Antecedente", "Fuente", "Valor"]
        filas = [encab]
        estilo_p = ParagraphStyle("celda", fontSize=7, leading=9)
        for a in alertas[:200]:
            filas.append([
                Paragraph(f"<b>{a.get('nivel','')}</b>", ParagraphStyle("n", fontSize=7, textColor=NIVEL_COLOR.get(a.get('nivel',''), GRIS))),
                Paragraph(f"{a.get('proveedor_secop','')}<br/>CC {a.get('documento','')}", estilo_p),
                Paragraph(f"{a.get('sancion','')}<br/>{a.get('tipo','')}", estilo_p),
                Paragraph(a.get("fuente","").replace("(","<br/>("), estilo_p),
                Paragraph(_fmt_cop(a.get("valor","")), estilo_p),
            ])
        ta = Table(filas, colWidths=[1.8*cm, 4.5*cm, 4.5*cm, 2.7*cm, 3*cm], repeatRows=1)
        ta.setStyle(TableStyle([
            ("BACKGROUND", (0,0), (-1,0), AZUL_OSCURO),
            ("TEXTCOLOR", (0,0), (-1,0), colors.white),
            ("FONTNAME", (0,0), (-1,0), "Helvetica-Bold"),
            ("FONTSIZE", (0,0), (-1,0), 8),
            ("GRID", (0,0), (-1,-1), 0.4, GRIS_CLARO),
            ("ROWBACKGROUNDS", (0,1), (-1,-1), [colors.white, GRIS_CLARO]),
            ("VALIGN", (0,0), (-1,-1), "TOP"),
            ("TOPPADDING", (0,0), (-1,-1), 3),
            ("BOTTOMPADDING", (0,0), (-1,-1), 3),
        ]))
        elems.append(ta)
        if len(alertas) > 200:
            elems.append(Spacer(1, 0.3*cm))
            elems.append(Paragraph(f"Se muestran las primeras 200 de {len(alertas)} alertas. El detalle completo est&aacute; en el archivo Excel.", styles["Pie"]))

    def _pie(canvas, doc_):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(GRIS)
        canvas.drawString(2*cm, 1.2*cm, f"Vigilancia Fiscal IA · XBOG SMART · Generado {datetime.now():%Y-%m-%d %H:%M}")
        canvas.drawRightString(19*cm, 1.2*cm, f"P\u00e1gina {doc_.page}")
        canvas.restoreState()

    doc.build(elems, onFirstPage=_pie, onLaterPages=_pie)
    buf.seek(0)
    return buf

def generar_excel(resultado):
    """Genera un Excel en memoria a partir del resultado del analisis."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Alertas"

    azul = PatternFill("solid", fgColor="0F172A")
    blanco = Font(color="FFFFFF", bold=True)
    borde = Border(*[Side(style="thin", color="E2E8F0")]*4)

    # Encabezado de resumen
    ws["A1"] = "VIGILANCIA FISCAL IA"
    ws["A1"].font = Font(size=14, bold=True, color="0F172A")
    ws["A2"] = f"Entidad: {resultado.get('entidad','')}"
    ws["A3"] = f"Fecha: {resultado.get('fecha_analisis','')}"
    ws["A4"] = f"Contratos: {resultado.get('total_contratos',0)} | Alertas: {resultado.get('total_alertas',0)} | Criticas: {resultado.get('criticas',0)} | Altas: {resultado.get('altas',0)}"

    encab = ["Nivel","Documento","Proveedor SECOP","Nombre fuente","Sancion","Tipo",
             "Fecha sancion","Fin inhabilidad","Vigente","Fuente","Valor contrato",
             "Estado","Fecha firma","Entidad","Objeto","Autoridad"]
    fila0 = 6
    for col, h in enumerate(encab, start=1):
        c = ws.cell(row=fila0, column=col, value=h)
        c.fill = azul; c.font = blanco; c.border = borde
        c.alignment = Alignment(horizontal="center", vertical="center")

    campos = ["nivel","documento","proveedor_secop","nombre_fuente","sancion","tipo",
              "fecha_sancion","fin_inhabilidad","vigente","fuente","valor",
              "estado","fecha_firma","entidad","objeto","autoridad"]
    for i, a in enumerate(resultado.get("alertas", []), start=fila0+1):
        for col, campo in enumerate(campos, start=1):
            ws.cell(row=i, column=col, value=a.get(campo, "")).border = borde

    anchos = [10,14,28,28,24,20,14,14,9,22,16,14,14,28,40,28]
    for col, w in enumerate(anchos, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = w
    ws.freeze_panes = f"A{fila0+1}"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
